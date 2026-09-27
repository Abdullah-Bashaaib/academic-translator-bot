"""
🎓 Academic Document Translator Bot — Final Edition
---------------------------------------------------
بوت تليجرام تفاعلي لترجمة المستندات الأكاديمية (PDF / PPTX)
مع اكتشاف تلقائي لموديل Gemini + واجهة أزرار وإيموجيات.
"""

import os
import asyncio
import logging
import tempfile
import time
from pathlib import Path
from dataclasses import dataclass

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

import arabic_reshaper
from bidi.algorithm import get_display

import google.generativeai as genai

from telegram import (
    Update,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    ReplyKeyboardMarkup,
    KeyboardButton,
    BotCommand,
)
from telegram.constants import ChatAction, ParseMode
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ContextTypes,
    filters,
)

# ==============================
# ⚙️ الإعدادات
# ==============================
logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)

TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
CHANNEL_ID = os.getenv("CHANNEL_ID")
CHANNEL_URL = os.getenv("CHANNEL_URL", "")
GEMINI_MODEL_ENV = os.getenv("GEMINI_MODEL", "")  # اختياري

MAX_FILE_SIZE_MB = 20
MAX_PAGES = 60

if not TELEGRAM_BOT_TOKEN or not GEMINI_API_KEY or not CHANNEL_ID:
    raise RuntimeError("❌ يجب ضبط: TELEGRAM_BOT_TOKEN, GEMINI_API_KEY, CHANNEL_ID")

# ==============================
# 🤖 تهيئة Gemini مع اكتشاف تلقائي للموديل
# ==============================
def _pick_gemini_model(preferred: str = "") -> str:
    """يحاول إيجاد موديل متاح من قائمة أولويات."""
    priority = [
        preferred,
        "gemini-2.5-flash",
        "gemini-2.5-flash-lite",
        "gemini-2.0-flash",
        "gemini-2.0-flash-001",
        "gemini-2.0-flash-lite",
        "gemini-flash-latest",
        "gemini-1.5-flash-latest",
        "gemini-pro-latest",
    ]
    priority = [m for m in priority if m]

    try:
        available = {
            m.name.replace("models/", "")
            for m in genai.list_models()
            if "generateContent" in m.supported_generation_methods
        }
        logger.info(f"📋 الموديلات المتاحة لحسابك: {sorted(available)}")

        for name in priority:
            if name in available:
                logger.info(f"✅ تم اختيار الموديل: {name}")
                return name

        if available:
            fallback = sorted(available)[0]
            logger.warning(f"⚠️ لا يوجد موديل من الأولويات، أستخدم: {fallback}")
            return fallback
    except Exception as e:
        logger.error(f"❌ فشل جلب قائمة الموديلات: {e}")

    return preferred or "gemini-2.5-flash"


genai.configure(api_key=GEMINI_API_KEY)
GEMINI_MODEL = _pick_gemini_model(GEMINI_MODEL_ENV)
gemini_model = genai.GenerativeModel(GEMINI_MODEL)

# ==============================
# 🗄️ إعدادات المستخدم
# ==============================
@dataclass
class UserPrefs:
    language: str = "ar"
    style: str = "academic"
    quality: str = "flash"
    show_progress: bool = True


USER_PREFS: dict[int, UserPrefs] = {}


def get_prefs(user_id: int) -> UserPrefs:
    if user_id not in USER_PREFS:
        USER_PREFS[user_id] = UserPrefs()
    return USER_PREFS[user_id]


# ==============================
# 🧠 محرك الترجمة
# ==============================
STYLE_PROMPTS = {
    "academic": (
        "أنت مترجم أكاديمي محترف. ترجم إلى العربية الفصحى بدقة عالية. "
        "حافظ على: المعادلات الرياضية، الرموز، الأكواد، المراجع، وأسماء الأعلام. "
        "لا تُضف شرحاً، أعد الترجمة فقط."
    ),
    "simple": (
        "ترجم النص التالي إلى العربية بأسلوب مبسّط وواضح لطالب جامعي. "
        "حافظ على المصطلحات العلمية بين قوسين. أعد الترجمة فقط."
    ),
    "literal": (
        "ترجم النص التالي ترجمة حرفية دقيقة كلمة بكلمة مع الحفاظ على البنية. "
        "أعد الترجمة فقط."
    ),
}


async def translate_text(text: str, style: str = "academic", retries: int = 4) -> str:
    """ترجمة موحّدة مع إعادة محاولة أسّية."""
    if not text or not text.strip():
        return ""

    system_prompt = STYLE_PROMPTS.get(style, STYLE_PROMPTS["academic"])
    full_prompt = f"{system_prompt}\n\n---\n{text}"

    delay = 2
    for attempt in range(1, retries + 1):
        try:
            response = await asyncio.to_thread(
                gemini_model.generate_content,
                full_prompt,
            )
            return (response.text or "").strip()
        except Exception as e:
            logger.warning(f"[Gemini] attempt {attempt}/{retries} failed: {e}")
            if attempt == retries:
                return f"[⚠️ فشل الترجمة: {str(e)[:100]}]"
            await asyncio.sleep(delay)
            delay *= 2
    return ""


def shape_arabic(text: str) -> str:
    if not text:
        return ""
    try:
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


# ==============================
# 🎨 لوحات المفاتيح
# ==============================
def main_menu_keyboard() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        [
            [KeyboardButton("📄 كيف أستخدم البوت؟"), KeyboardButton("⚙️ الإعدادات")],
            [KeyboardButton("ℹ️ معلومات البوت"), KeyboardButton("🆘 مساعدة")],
        ],
        resize_keyboard=True,
        input_field_placeholder="أرسل ملف PDF أو PPTX لترجمته...",
    )


def subscription_keyboard() -> InlineKeyboardMarkup:
    buttons = []
    if CHANNEL_URL:
        buttons.append([InlineKeyboardButton("📢 اشترك الآن", url=CHANNEL_URL)])
    buttons.append([InlineKeyboardButton("✅ تحققت من الاشتراك", callback_data="check_sub")])
    return InlineKeyboardMarkup(buttons)


def settings_keyboard(prefs: UserPrefs) -> InlineKeyboardMarkup:
    style_map = {"academic": "🎓 أكاديمي", "simple": "💡 مبسّط", "literal": "📖 حرفي"}

    return InlineKeyboardMarkup([
        [InlineKeyboardButton("🎨 أسلوب الترجمة", callback_data="noop")],
        [
            InlineKeyboardButton(
                ("✅ " if prefs.style == k else "") + v,
                callback_data=f"set_style_{k}",
            )
            for k, v in style_map.items()
        ],
        [InlineKeyboardButton("📊 شريط التقدم", callback_data="noop")],
        [
            InlineKeyboardButton(
                ("✅ مفعّل" if prefs.show_progress else "❌ معطّل"),
                callback_data="toggle_progress",
            ),
        ],
        [InlineKeyboardButton("🔙 إغلاق", callback_data="close_settings")],
    ])


def after_result_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup([
        [
            InlineKeyboardButton("🔄 ترجمة ملف آخر", callback_data="new_file"),
            InlineKeyboardButton("⚙️ تغيير الإعدادات", callback_data="open_settings"),
        ],
        [InlineKeyboardButton("⭐ قيّم تجربتك", callback_data="rate_bot")],
    ])


# ==============================
# 🔐 التحقق من الاشتراك
# ==============================
async def is_subscribed(update: Update, context: ContextTypes.DEFAULT_TYPE) -> bool:
    user_id = update.effective_user.id
    try:
        member = await context.bot.get_chat_member(chat_id=CHANNEL_ID, user_id=user_id)
        return member.status in ("member", "administrator", "creator")
    except Exception as e:
        logger.error(f"get_chat_member failed: {e}")
        return False


async def send_subscription_prompt(update: Update):
    text = (
        "🔒 *الاشتراك مطلوب للاستخدام*\n\n"
        "من فضلك اشترك في قناتنا أولاً لدعم البوت،\n"
        "ثم اضغط *«تحققت من الاشتراك»* ✅\n\n"
        "_نعتذر عن معالجة ملفك قبل إتمام الاشتراك_ 🙏"
    )
    markup = subscription_keyboard()
    if update.callback_query:
        await update.callback_query.edit_message_text(
            text, parse_mode=ParseMode.MARKDOWN, reply_markup=markup
        )
    else:
        await update.message.reply_text(
            text, parse_mode=ParseMode.MARKDOWN, reply_markup=markup
        )


# ==============================
# 📄 معالجة PDF
# ==============================
async def process_pdf(input_path: str, output_path: str, style: str, progress_cb=None) -> int:
    src_doc = fitz.open(input_path)
    out_doc = fitz.open()
    total = len(src_doc)

    if total > MAX_PAGES:
        raise ValueError(f"الحد الأقصى {MAX_PAGES} صفحة (ملفك: {total})")

    for i, page in enumerate(src_doc):
        out_doc.insert_pdf(src_doc, from_page=i, to_page=i)

        raw_text = page.get_text("text").strip()
        translated = await translate_text(raw_text, style=style) if raw_text else ""
        shaped = shape_arabic(translated)

        rect = page.rect
        new_page = out_doc.new_page(width=rect.width, height=rect.height)

        if shaped:
            try:
                new_page.insert_textbox(
                    fitz.Rect(40, 40, rect.width - 40, rect.height - 40),
                    shaped,
                    fontsize=13,
                    fontname="helv",
                    align=fitz.TEXT_ALIGN_RIGHT,
                )
            except Exception as e:
                new_page.insert_text((40, 60), f"[خطأ عرض: {e}]", fontsize=11)
        else:
            new_page.insert_text((40, 60), "[صفحة بدون نص]", fontsize=11)

        if progress_cb:
            await progress_cb(i + 1, total)

    out_doc.save(output_path, garbage=4, deflate=True)
    out_doc.close()
    src_doc.close()
    return total


# ==============================
# 📊 معالجة PPTX
# ==============================
async def process_pptx(input_path: str, output_path: str, style: str, progress_cb=None) -> int:
    prs = Presentation(input_path)
    total = len(prs.slides)
    if total > MAX_PAGES:
        raise ValueError(f"الحد الأقصى {MAX_PAGES} شريحة (ملفك: {total})")

    slide_texts = []
    for slide in prs.slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        parts.append(line.strip())
        slide_texts.append("\n".join(parts))

    for idx, text in enumerate(slide_texts):
        translated = await translate_text(text, style=style) if text else ""
        blank_layout = prs.slide_layouts[6]
        new_slide = prs.slides.add_slide(blank_layout)

        xml_slides = prs.slides._sldIdLst  # noqa
        slides_list = list(xml_slides)
        target_pos = (idx * 2) + 1
        if target_pos < len(slides_list):
            xml_slides.remove(slides_list[-1])
            xml_slides.insert(target_pos, slides_list[-1])

        tx_box = new_slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5),
            prs.slide_width - Inches(1), prs.slide_height - Inches(1),
        )
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.clear()

        lines = translated.split("\n") if translated else ["[لا يوجد نص]"]
        for j, line in enumerate(lines):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = line
            run.font.size = Pt(18)
            run.font.name = "Arial"

        if progress_cb:
            await progress_cb(idx + 1, total)

    prs.save(output_path)
    return total


# ==============================
# 🎬 الأوامر
# ==============================
async def start_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    if not await is_subscribed(update, context):
        await send_subscription_prompt(update)
        return

    text = (
        f"👋 *أهلاً بك يا {user.first_name}!*\n\n"
        "🎓 *بوت الترجمة الأكاديمية*\n"
        "أترجم مستنداتك (PDF • PPTX) إلى العربية بأسلوب أكاديمي دقيق.\n\n"
        "📌 *كيف يعمل؟*\n"
        "أرسل لي ملفك وسأعيده لك بصيغة:\n"
        "_الصفحة الأصلية ➡️ الترجمة ➡️ الأصلية ➡️ الترجمة ..._\n\n"
        f"⚠️ الحد الأقصى: *{MAX_PAGES}* صفحة • *{MAX_FILE_SIZE_MB}* MB\n\n"
        "✨ *ابدأ الآن:* أرسل ملفك مباشرة 👇"
    )
    await update.message.reply_text(
        text, parse_mode=ParseMode.MARKDOWN, reply_markup=main_menu_keyboard()
    )


async def help_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = (
        "🆘 *مركز المساعدة*\n\n"
        "🔹 *الصيغ المدعومة:*\n"
        "• PDF (مستندات)\n"
        "• PPTX (عروض تقديمية)\n\n"
        "🔹 *طريقة العمل:*\n"
        "1️⃣ أرسل الملف\n"
        "2️⃣ انتظر الترجمة (قد تأخذ دقائق)\n"
        "3️⃣ استلم الملف المترجم مباشرة\n\n"
        "🔹 *الأوامر:*\n"
        "• /start — البداية\n"
        "• /help — المساعدة\n"
        "• /settings — الإعدادات\n"
        "• /about — عن البوت\n\n"
        "💡 *نصيحة:* كلما كان الملف أصغر، كانت الترجمة أسرع!"
    )
    await update.message.reply_text(
        text, parse_mode=ParseMode.MARKDOWN, reply_markup=main_menu_keyboard()
    )


async def about_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = (
        "ℹ️ *عن البوت*\n\n"
        "🤖 *الاسم:* Academic Translator Bot\n"
        f"🧠 *المحرك:* Google Gemini (`{GEMINI_MODEL}`)\n"
        "🌍 *اللغات:* من الإنجليزية إلى العربية\n"
        "🎯 *التخصص:* المستندات الأكاديمية والعلمية\n\n"
        "🔒 *الخصوصية:*\n"
        "• لا نحفظ ملفاتك على السيرفر\n"
        "• تُحذف تلقائياً بعد الترجمة\n\n"
        "💚 *تطوير:* مجتمعي مفتوح"
    )
    await update.message.reply_text(
        text, parse_mode=ParseMode.MARKDOWN, reply_markup=main_menu_keyboard()
    )


async def settings_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await is_subscribed(update, context):
        await send_subscription_prompt(update)
        return

    prefs = get_prefs(update.effective_user.id)
    text = "⚙️ *الإعدادات*\n\nاضبط تفضيلاتك لترجمة تناسب احتياجاتك 👇"
    await update.message.reply_text(
        text, parse_mode=ParseMode.MARKDOWN, reply_markup=settings_keyboard(prefs)
    )


# ==============================
# 🎬 Callbacks
# ==============================
async def on_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    data = query.data
    user_id = update.effective_user.id
    prefs = get_prefs(user_id)

    if data == "check_sub":
        await query.answer()
        if await is_subscribed(update, context):
            await query.edit_message_text(
                "✅ *تم التحقق من اشتراكك بنجاح!*\n\n"
                "🎉 مرحباً بك، يمكنك الآن إرسال ملفك.\n"
                "أرسل /start للبدء.",
                parse_mode=ParseMode.MARKDOWN,
            )
        else:
            await query.answer("❌ لم يتم العثور على اشتراكك بعد.", show_alert=True)
        return

    if data == "open_settings":
        await query.answer()
        await query.edit_message_text(
            "⚙️ *الإعدادات*\n\nاضبط تفضيلاتك 👇",
            parse_mode=ParseMode.MARKDOWN,
            reply_markup=settings_keyboard(prefs),
        )
        return

    if data == "noop":
        await query.answer()
        return

    if data.startswith("set_style_"):
        prefs.style = data.replace("set_style_", "")
        await query.answer("✅ تم التحديث")
        await query.edit_message_reply_markup(settings_keyboard(prefs))
        return

    if data == "toggle_progress":
        prefs.show_progress = not prefs.show_progress
        await query.answer("✅ تم التحديث")
        await query.edit_message_reply_markup(settings_keyboard(prefs))
        return

    if data == "close_settings":
        await query.answer("تم الإغلاق")
        await query.edit_message_text("⚙️ *تم حفظ إعداداتك* ✅", parse_mode=ParseMode.MARKDOWN)
        return

    if data == "new_file":
        await query.answer()
        await query.edit_message_text(
            "📤 *أرسل ملفك الجديد الآن*\n\nPDF أو PPTX — وسأبدأ الترجمة فوراً.",
            parse_mode=ParseMode.MARKDOWN,
        )
        return

    if data == "rate_bot":
        await query.answer("شكراً لك! 💚", show_alert=True)
        return


# ==============================
# 🎬 معالج الأزرار النصية
# ==============================
async def on_menu_button(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text
    if text == "📄 كيف أستخدم البوت؟":
        await help_cmd(update, context)
    elif text == "⚙️ الإعدادات":
        await settings_cmd(update, context)
    elif text == "ℹ️ معلومات البوت":
        await about_cmd(update, context)
    elif text == "🆘 مساعدة":
        await help_cmd(update, context)


# ==============================
# 📥 معالج الملفات
# ==============================
def format_progress(current: int, total: int) -> str:
    percent = current / total if total else 0
    filled = int(percent * 10)
    bar = "▰" * filled + "▱" * (10 - filled)
    return f"{bar} {int(percent * 100)}%"


async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await is_subscribed(update, context):
        await send_subscription_prompt(update)
        return

    doc = update.message.document
    if not doc:
        return

    user_id = update.effective_user.id
    prefs = get_prefs(user_id)

    filename = doc.file_name or "file"
    ext = Path(filename).suffix.lower()

    if ext not in (".pdf", ".pptx"):
        await update.message.reply_text(
            "⚠️ *صيغة غير مدعومة*\n\n"
            "الصيغ المتاحة: *PDF* • *PPTX*\n"
            "أرسل ملفاً بصيغة صحيحة 🙏",
            parse_mode=ParseMode.MARKDOWN,
        )
        return

    size_mb = (doc.file_size or 0) / (1024 * 1024)
    if size_mb > MAX_FILE_SIZE_MB:
        await update.message.reply_text(
            f"⚠️ *الملف كبير جداً*\n\n"
            f"📦 حجم ملفك: `{size_mb:.1f} MB`\n"
            f"🎯 الحد الأقصى: `{MAX_FILE_SIZE_MB} MB`\n\n"
            "قسّم الملف وأعد الإرسال 🙏",
            parse_mode=ParseMode.MARKDOWN,
        )
        return

    icon = "📕" if ext == ".pdf" else "📊"
    style_map = {"academic": "🎓 أكاديمي", "simple": "💡 مبسّط", "literal": "📖 حرفي"}

    status_msg = await update.message.reply_text(
        f"{icon} *استلمت ملفك!*\n\n"
        f"📎 *الاسم:* `{filename}`\n"
        f"📦 *الحجم:* `{size_mb:.2f} MB`\n"
        f"🎨 *الأسلوب:* {style_map.get(prefs.style, '🎓 أكاديمي')}\n\n"
        "⏳ *جارٍ التحميل...*",
        parse_mode=ParseMode.MARKDOWN,
    )

    await context.bot.send_chat_action(update.effective_chat.id, ChatAction.TYPING)

    tmp_in = tmp_out = None
    start_time = time.time()

    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tf:
            tmp_in = tf.name

        tg_file = await context.bot.get_file(doc.file_id)
        await tg_file.download_to_drive(tmp_in)

        last_update = [0]

        async def progress_cb(current: int, total: int):
            if not prefs.show_progress:
                return
            now = time.time()
            if now - last_update[0] < 3 and current != total:
                return
            last_update[0] = now

            bar = format_progress(current, total)
            try:
                await status_msg.edit_text(
                    f"🧠 *جارٍ الترجمة...*\n\n"
                    f"{bar}\n\n"
                    f"📄 الصفحة: `{current} / {total}`\n"
                    f"⏱️ مضى: `{int(now - start_time)}` ثانية",
                    parse_mode=ParseMode.MARKDOWN,
                )
            except Exception:
                pass

        await status_msg.edit_text(
            "🧠 *جارٍ الترجمة...*\n\n"
            f"{format_progress(0, 1)}\n\n"
            "_يستغرق الأمر دقائق حسب حجم الملف_",
            parse_mode=ParseMode.MARKDOWN,
        )

        tmp_out = tmp_in.replace(ext, f"_translated{ext}")

        if ext == ".pdf":
            pages = await process_pdf(tmp_in, tmp_out, prefs.style, progress_cb)
        else:
            pages = await process_pptx(tmp_in, tmp_out, prefs.style, progress_cb)

        elapsed = int(time.time() - start_time)
        await status_msg.edit_text("📤 *جارٍ إرسال الملف...*", parse_mode=ParseMode.MARKDOWN)

        out_name = f"translated_{Path(filename).stem}{ext}"
        with open(tmp_out, "rb") as f:
            await update.message.reply_document(
                document=f,
                filename=out_name,
                caption=(
                    f"✅ *اكتملت الترجمة بنجاح!*\n\n"
                    f"📄 *الصفحات:* `{pages}`\n"
                    f"⏱️ *الوقت:* `{elapsed}` ثانية\n"
                    f"🎨 *الأسلوب:* {style_map.get(prefs.style)}\n"
                    f"📎 `{out_name}`"
                ),
                parse_mode=ParseMode.MARKDOWN,
                reply_markup=after_result_keyboard(),
            )

        await status_msg.delete()

    except ValueError as e:
        await status_msg.edit_text(
            f"⚠️ *تعذّر معالجة الملف*\n\n`{e}`",
            parse_mode=ParseMode.MARKDOWN,
        )
    except Exception as e:
        logger.exception("Processing failed")
        await status_msg.edit_text(
            "❌ *حدث خطأ غير متوقع*\n\n"
            "حاول مرة أخرى، وإن تكرر الخطأ:\n"
            f"`{str(e)[:200]}`",
            parse_mode=ParseMode.MARKDOWN,
        )
    finally:
        for p in (tmp_in, tmp_out):
            try:
                if p and os.path.exists(p):
                    os.remove(p)
            except Exception as e:
                logger.warning(f"Cleanup failed: {e}")


async def error_handler(update: object, context: ContextTypes.DEFAULT_TYPE):
    logger.error("Exception:", exc_info=context.error)


# ==============================
# 🚀 نقطة الانطلاق
# ==============================
async def post_init(app: Application):
    await app.bot.set_my_commands([
        BotCommand("start", "🏠 البداية"),
        BotCommand("help", "🆘 المساعدة"),
        BotCommand("settings", "⚙️ الإعدادات"),
        BotCommand("about", "ℹ️ عن البوت"),
    ])


def main():
    app = Application.builder().token(TELEGRAM_BOT_TOKEN).post_init(post_init).build()

    app.add_handler(CommandHandler("start", start_cmd))
    app.add_handler(CommandHandler("help", help_cmd))
    app.add_handler(CommandHandler("settings", settings_cmd))
    app.add_handler(CommandHandler("about", about_cmd))
    app.add_handler(CallbackQueryHandler(on_callback))
    app.add_handler(MessageHandler(
        filters.Regex("^(📄 كيف أستخدم البوت؟|⚙️ الإعدادات|ℹ️ معلومات البوت|🆘 مساعدة)$"),
        on_menu_button,
    ))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    app.add_error_handler(error_handler)

    logger.info("🤖 Bot is running...")
    app.run_polling(allowed_updates=Update.ALL_TYPES, drop_pending_updates=True)


if __name__ == "__main__":
    main()
